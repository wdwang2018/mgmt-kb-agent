"""课前预热物料生成测试（P7）：brochure 抽取 + docx/pptx/pdf 渲染 + 端点全链路。

含变异验证：改坏实现（忽略 theme_hint / 不写标题）对应测试必须失败，
确保测试不是空转（项目铁律）。
"""

from __future__ import annotations

import io
import json
import zipfile

from app.domain.models import CoursePackage
from app.services.course_materials import (build_brochure, build_materials_zip,
                                          render_docx, render_pptx, render_pdf)

# 紧凑但完整的样本包：覆盖 need/gap/objectives/modules/content/citations
SAMPLE = {
    "run_id": "run-test-sample",
    "need": {
        "audience": "新任经理", "audience_level": 2, "headcount": 24,
        "duration_minutes": 360,
        "business_pain": "新经理仍埋头做事，下属得不到成长",
        "desired_behavior": "会授权、会给反馈、坚持 1on1",
        "industry": "互联网", "theme_hint": "新任管理者转型",
    },
    "gap": {
        "targets": [
            {"domain": "TEAM", "target_level": 3, "gap": "L2→L3",
             "rationale": "痛点映射"},
            {"domain": "COACH", "target_level": 3, "gap": "L2→L3",
             "rationale": "痛点映射"},
        ],
        "summary": "核心差距在团队领导与教练辅导。", "dilemmas": ["D01", "D10"],
    },
    "objectives": [
        {"id": "OBJ-1", "text": "能够在真实情境中组建跨职能团队并处理冲突", "bloom": 3,
         "domain": "TEAM", "target_level": 3},
        {"id": "OBJ-2", "text": "能够用 GROW 模型辅导下属", "bloom": 3,
         "domain": "COACH", "target_level": 3},
    ],
    "modules": [
        {"id": "M1", "title": "开场：转身之痛", "minutes": 30,
         "objective_ids": ["OBJ-1"], "bucket_mix": {"formal": 20, "social": 0, "practice": 10},
         "focus_domains": ["TEAM"], "focus_dilemmas": ["D01"]},
        {"id": "M2", "title": "授权：五级阶梯", "minutes": 60,
         "objective_ids": ["OBJ-1", "OBJ-2"],
         "bucket_mix": {"formal": 17, "social": 13, "practice": 30},
         "focus_domains": ["TEAM", "COACH"], "focus_dilemmas": ["D02"]},
    ],
    "content": {
        "M1": [{"text": "管理是计划/组织/指挥/协调/控制五大职能", "unit_id": "KU-CLASSIC-002",
                "citation": "Fayol《工业管理与一般管理》", "evidence_level": "C",
                "misconceptions": ["管理就是发号施令"]}],
        "M2": [{"text": "X-Y 理论：人性假设决定管理方式", "unit_id": "KU-BEHAV-002",
                "citation": "McGregor《企业的人性面》", "evidence_level": "C",
                "misconceptions": ["Y 理论就是放任"]}],
    },
    "citations": ["Fayol《工业管理与一般管理》", "McGregor《企业的人性面》", "Mintzberg《管理工作的本质》"],
}


def _pkg():
    return CoursePackage.model_validate(SAMPLE)


def test_build_brochure_uses_theme_hint():
    """标题应优先采用 theme_hint（变异点：若实现忽略 theme_hint 必失败）。"""
    b = build_brochure(_pkg())
    assert b.title == "新任管理者转型", b.title
    assert "互联网" in b.subtitle
    assert "24 人小班" in b.meta_line
    assert "1 天" in b.meta_line
    # 能力域中文名应被映射进亮点
    joined = "｜".join(b.highlights)
    assert "团队领导" in joined and "教练辅导" in joined
    # 引用数与课时配比
    assert b.citations_count == 3
    assert len(b.outline) == 2
    assert len(b.takeaways) == 2


def test_build_brochure_no_theme_fallback():
    """无 theme_hint 时退化为「受众+管理转身」，不崩。"""
    d = json.loads(json.dumps(SAMPLE))
    d["need"].pop("theme_hint")
    b = build_brochure(CoursePackage.model_validate(d))
    assert b.title == "新任经理管理转身"


def test_renderers_produce_valid_files():
    """三端渲染产物应为合法文件，且 docx/pptx 必须含标题（内容校验，非空转）。"""
    b = build_brochure(_pkg())
    docx = render_docx(b)
    pptx = render_pptx(b)
    pdf = render_pdf(b)
    assert docx[:4] == b"PK\x03\x04"
    assert pptx[:2] == b"PK"
    assert pdf[:5] == b"%PDF-"
    assert len(docx) > 3000 and len(pptx) > 3000 and len(pdf) > 1000

    # docx 重新打开，正文必须包含标题与亮点关键字
    from docx import Document
    doc = Document(io.BytesIO(docx))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "新任管理者转型" in text, "docx 未写入标题"
    assert "课程亮点" in text
    assert "Fayol" in text, "docx 未写入理论要点"

    # pptx 重新打开，至少一页含标题
    from pptx import Presentation
    prs = Presentation(io.BytesIO(pptx))
    all_text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text.append(shape.text_frame.text)
    blob = "\n".join(all_text)
    assert "新任管理者转型" in blob, "pptx 未写入标题"
    assert "课程大纲" in blob


def test_build_materials_zip():
    """zip 应包含 docx/pptx/pdf 三件套且体积合理。"""
    zb = build_materials_zip(build_brochure(_pkg()), "run-test-sample")
    z = zipfile.ZipFile(io.BytesIO(zb))
    names = z.namelist()
    assert any(n.endswith(".docx") for n in names)
    assert any(n.endswith(".pptx") for n in names)
    assert any(n.endswith(".pdf") for n in names)
    for n in names:
        assert z.getinfo(n).file_size > 1000


def test_endpoint_full_flow_and_404():
    """端点全链路：跑通流水线 → 取物料 zip；不存在的 run 应 404。"""
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from fastapi.testclient import TestClient

    with TestClient(__import__("app.api.main", fromlist=["app"]).app) as client:
        NEED = {
            "audience": "新任经理", "audience_level": 2, "headcount": 24,
            "duration_minutes": 360,
            "business_pain": "新经理仍埋头做事，下属得不到成长",
            "desired_behavior": "会授权、会给反馈、坚持 1on1",
            "industry": "互联网", "theme_hint": "新任管理者转型",
        }
        r = client.post("/api/v1/runs", json=NEED)
        assert r.status_code == 200
        run_id = r.json()["run_id"]
        r2 = client.post(f"/api/v1/runs/{run_id}/resume",
                         json={"decision": {"action": "approve"}})
        assert r2.status_code == 200 and r2.json()["status"] == "completed"

        m = client.get(f"/api/v1/runs/{run_id}/materials")
        assert m.status_code == 200
        assert m.headers["content-type"].startswith("application/zip")
        assert "attachment" in m.headers.get("content-disposition", "")
        z = zipfile.ZipFile(io.BytesIO(m.content))
        assert sum(1 for n in z.namelist() if n.endswith((".docx", ".pptx", ".pdf"))) == 3

        # 不存在的 run
        r3 = client.get("/api/v1/runs/run-does-not-exist/materials")
        assert r3.status_code == 404

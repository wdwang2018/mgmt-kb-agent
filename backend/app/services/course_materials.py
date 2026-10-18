"""课前预热物料生成服务（P7）。

把一份已生成的课程包（CoursePackage）转换成「授课前预热 / 招生宣传」用的
一套物料：Word（宣传介绍）、PPT（宣讲/预热幻灯片）、PDF（可打印宣传页），
打包成 zip 一次性返回。

设计约束（延续项目铁律）：
- 纯数据转换，**不调用任何 LLM**——课程包已落盘，离线 Mock 模式下一键可出。
- 学术克制风格：墨蓝主色 + 暖中性分隔线 + 衬线标题，无 emoji、无浮夸用语。
- 内容全部由课程包字段派生，不编造；引用遵循「要点化转述、可溯源」的语义边界。
- 中文 CJK 字体：docx/pptx 显式设置东亚字体；pdf 用 reportlab 内置 STSong-Light CID 字体（无需字体文件）。
"""

from __future__ import annotations

import io
import json
import zipfile
from dataclasses import dataclass, field
from typing import Any

from ..domain.constants import COMPETENCY_DOMAINS
from ..domain.models import CoursePackage

# —— 学术克制调色板 ——
INK = "2B4C6F"                       # 墨蓝主色
INK_RGB = (0x2B, 0x4C, 0x6F)
ACCENT = "8A6D3B"                    # 暖金（仅用于极少量强调，克制）
TEXT = "222222"
MUTED = "5A5A5A"
RULE = "D9D2C5"                      # 暖中性分隔线
PAPER = "FBF9F4"                     # 纸感底色

CN_FONT = "微软雅黑"                 # docx/pptx 东亚字体
CN_SERIF = "宋体"                    # 衬线标题（学术克制）
PDF_FONT = "STSong-Light"            # reportlab CID 中文字体（内置）


# ---------------------------------------------------------------------------
# 数据模型：从课程包抽取的「预热物料」结构（三端渲染共用单一数据源）
# ---------------------------------------------------------------------------
@dataclass
class ModuleCard:
    idx: int
    title: str
    minutes: int
    points: list[str] = field(default_factory=list)


@dataclass
class Brochure:
    title: str
    subtitle: str
    audience: str
    meta_line: str
    pain: str
    outcome: str
    highlights: list[str]
    outline: list[ModuleCard]
    takeaways: list[str]
    citation_note: str
    cta: str
    citations_count: int


def _dur_label(minutes: int) -> str:
    """把分钟数转成友好时长描述。"""
    if minutes <= 0:
        return "时长待定"
    h = minutes / 60
    if h >= 6:
        return f"1 天（{int(round(h))} 小时）"
    if h >= 3:
        return f"半天至 1 天（约 {h:.0f} 小时）"
    if h >= 1:
        return f"约 {h:.0f} 小时"
    return f"{minutes} 分钟"


def _clean(s: str) -> str:
    return (s or "").replace("\n", " ").strip()


def build_brochure(pkg: CoursePackage | dict) -> Brochure:
    """从课程包抽取预热物料结构。接受 CoursePackage 或原始 dict。"""
    if isinstance(pkg, dict):
        pkg = CoursePackage.model_validate(pkg)
    need = pkg.need
    gap = pkg.gap

    theme = _clean(need.theme_hint)
    title = theme if theme else f"{_clean(need.audience)}管理转身"
    subtitle = f"面向{_clean(need.industry)}·{_clean(need.audience)}的研修 ｜ 基于管理经典与实证"

    meta = (f"时长：{_dur_label(need.duration_minutes)}"
            f" ｜ 班型：{need.headcount} 人小班"
            f" ｜ 形式：理论讲授 + 小组研讨 + 实战演练")

    # 能力域中文名（单一数据源：constants.COMPETENCY_DOMAINS）
    domains: list[str] = []
    for t in (gap.targets or []):
        if t.domain not in domains:
            domains.append(t.domain)
    domain_names = [COMPETENCY_DOMAINS.get(d, {}).get("name", d) for d in domains]

    # 70-20-10 配比聚合
    agg = {"formal": 0, "social": 0, "practice": 0}
    for m in pkg.modules:
        bm = m.bucket_mix or {}
        for k in agg:
            agg[k] += int(bm.get(k, 0) or 0)
    total = sum(agg.values()) or 1
    pct = lambda k: round(agg[k] / total * 100)

    highlights: list[str] = []
    if _clean(need.business_pain):
        highlights.append(f"聚焦真实业务痛点：{_clean(need.business_pain)}")
    if domain_names:
        highlights.append("系统突破 " + "、".join(domain_names) + " 等核心能力域")
    highlights.append(
        f"理论有据：引用 {len(pkg.citations)} 部管理经典与实证研究，要点化转述、每条可溯源")
    highlights.append(
        f"混合式设计：实战演练 {pct('practice')}% ＋ 小组研讨 {pct('social')}% ＋ 理论讲授 {pct('formal')}%（贴合 70-20-10 成人学习配比）")
    if _clean(need.desired_behavior):
        highlights.append(f"带得走的行为改变：{_clean(need.desired_behavior)}")

    outline: list[ModuleCard] = []
    for i, m in enumerate(pkg.modules, 1):
        pts = (pkg.content or {}).get(m.id, []) or []
        points = [_clean(p.text) for p in pts][:2]
        outline.append(ModuleCard(i, _clean(m.title), m.minutes, points))

    takeaways = [_clean(o.text) for o in pkg.objectives]

    top = (pkg.citations or [])[:2]
    if len(top) >= 2:
        cite_head = f"（如 {_clean(top[0])}、{_clean(top[1])} 等"
    elif top:
        cite_head = f"（如 {_clean(top[0])}"
    else:
        cite_head = ""
    citation_note = (f"课程内容均基于管理学经典与实证研究{cite_head}共 {len(pkg.citations)} 部），"
                    "要点化转述、可溯源，不代表原文逐字引用。")

    cta = ("报名与咨询：请填写下方信息或扫描二维码（占位）。开课时间、地点与费用以正式邀请函为准；"
           "团体定制与企业内训可联系课程顾问。")

    return Brochure(
        title=title, subtitle=subtitle, audience=_clean(need.audience),
        meta_line=meta, pain=_clean(need.business_pain),
        outcome=_clean(need.desired_behavior), highlights=highlights,
        outline=outline, takeaways=takeaways, citation_note=citation_note,
        cta=cta, citations_count=len(pkg.citations),
    )


# ---------------------------------------------------------------------------
# DOCX 渲染（宣传介绍）
# ---------------------------------------------------------------------------
def _docx_cjk(run, font: str = CN_FONT) -> None:
    """显式设置东亚字体，避免中文回退成方块。"""
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(__import__("docx").oxml.ns.qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(__import__("docx").oxml.ns.qn("w:eastAsia"), font)
    rfonts.set(__import__("docx").oxml.ns.qn("w:ascii"), font)
    rfonts.set(__import__("docx").oxml.ns.qn("w:hAnsi"), font)


def render_docx(b: Brochure) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    # 基础字体
    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(__import__("docx").oxml.ns.qn("w:eastAsia"), CN_FONT)

    ink = RGBColor(*INK_RGB)
    muted = RGBColor(0x5A, 0x5A, 0x5A)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = ink
        _docx_cjk(r, CN_SERIF)
        # 标题下细线
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        pPr = p2._p.get_or_add_pPr()
        from docx.oxml import OxmlElement
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(__import__("docx").oxml.ns.qn("w:val"), "single")
        bottom.set(__import__("docx").oxml.ns.qn("w:sz"), "6")
        bottom.set(__import__("docx").oxml.ns.qn("w:space"), "1")
        bottom.set(__import__("docx").oxml.ns.qn("w:color"), RULE)
        pbdr.append(bottom)
        pPr.append(pbdr)
        p2.paragraph_format.space_after = Pt(4)

    def body(text: str, color=None, bold=False, size=10.5) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        _docx_cjk(r)
        return p

    def bullet(text: str) -> None:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(text)
        r.font.size = Pt(10.5)
        _docx_cjk(r)

    # 封面标题
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run(b.title)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = ink
    _docx_cjk(r, CN_SERIF)

    s = doc.add_paragraph()
    rs = s.add_run(b.subtitle)
    rs.font.size = Pt(11)
    rs.font.color.rgb = muted
    _docx_cjk(rs)

    meta = doc.add_paragraph()
    rm = meta.add_run(b.meta_line)
    rm.font.size = Pt(10)
    rm.font.color.rgb = muted
    _docx_cjk(rm)

    # 课程亮点
    heading("课程亮点")
    for h in b.highlights:
        bullet(h)

    # 适合谁 / 为什么学
    heading("适合谁听")
    body(f"本课程面向：{b.audience}。", bold=True)
    if b.pain:
        body("你将面对的问题：" + b.pain)
    if b.outcome:
        body("学完你将能够：" + b.outcome)

    # 课程大纲
    heading("课程大纲")
    for m in b.outline:
        p = doc.add_paragraph()
        rt = p.add_run(f"模块{m.idx}　{m.title}（{m.minutes} 分钟）")
        rt.bold = True
        rt.font.size = Pt(11)
        rt.font.color.rgb = ink
        _docx_cjk(rt)
        for pt in m.points:
            bullet(pt)

    # 学习收获
    heading("学习收获")
    for tk in b.takeaways:
        bullet(tk)

    # 内容依据
    heading("内容依据")
    body(b.citation_note, color=muted, size=9.5)

    # 报名咨询
    heading("报名与咨询")
    body(b.cta)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PPTX 渲染（宣讲/预热幻灯片）
# ---------------------------------------------------------------------------
def _pptx_ea(run, font: str = CN_FONT) -> None:
    """设置东亚字体（a:ea），避免中文回退成方块。"""
    from pptx.oxml.ns import qn
    from pptx.oxml import parse_xml
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    # 移除已有 ea 再添加，避免重复
    for old in rpr.findall(qn("a:ea")):
        rpr.remove(old)
    rpr.append(parse_xml(f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'))


def _pptx_slide(prs, title: str):
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    # 顶部墨蓝标题条
    bar = slide.shapes.add_shape(
        1, Inches(0), Inches(0), prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*INK_RGB)
    bar.line.fill.background()
    # 标题
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), prs.slide_width - Inches(1), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = RGBColor(*INK_RGB)
    _pptx_ea(r, CN_SERIF)
    return slide


def _pptx_body(slide, prs, lines, top_inches=1.4, size=16):
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    box = slide.shapes.add_textbox(Inches(0.6), Inches(top_inches),
                                   prs.slide_width - Inches(1.2),
                                   prs.slide_height - Inches(top_inches + 0.4))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(8)
        p.alignment = PP_ALIGN.LEFT
        if isinstance(ln, tuple):
            prefix, text = ln
            r0 = p.add_run()
            r0.text = prefix
            r0.font.size = Pt(size)
            r0.font.bold = True
            r0.font.color.rgb = RGBColor(*INK_RGB)
            _pptx_ea(r0)
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            _pptx_ea(r)
        else:
            r = p.add_run()
            r.text = "• " + ln
            r.font.size = Pt(size)
            _pptx_ea(r)
    return box


def render_pptx(b: Brochure) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1) 封面
    cover = _pptx_slide(prs, b.title)
    _pptx_body(cover, prs, [b.subtitle, b.meta_line, f"面向对象：{b.audience}"],
               top_inches=1.6, size=18)
    cover.shapes.add_textbox(Inches(0.6), Inches(6.4), prs.slide_width - Inches(1.2), Inches(0.6))

    # 2) 课程亮点
    hl = _pptx_slide(prs, "课程亮点")
    _pptx_body(hl, prs, b.highlights, size=18)

    # 3) 适合谁 / 为什么学
    why = _pptx_slide(prs, "适合谁听 · 为什么学")
    lines = [f"面向对象：{b.audience}"]
    if b.pain:
        lines.append(("问题　", b.pain))
    if b.outcome:
        lines.append(("目标　", b.outcome))
    _pptx_body(why, prs, lines, size=18)

    # 4-5) 课程大纲（每页 4 个模块）
    chunk = 4
    for ci in range(0, len(b.outline), chunk):
        sl = _pptx_slide(prs, "课程大纲" if ci == 0 else "课程大纲（续）")
        lines = []
        for m in b.outline[ci:ci + chunk]:
            pts = ("：" + "；".join(m.points)) if m.points else ""
            lines.append((f"模块{m.idx}（{m.minutes}′）　", m.title + pts))
        _pptx_body(sl, prs, lines, size=15)

    # 6) 学习收获
    tk = _pptx_slide(prs, "学习收获")
    _pptx_body(tk, prs, b.takeaways, size=17)

    # 7) 内容依据 + 报名
    end = _pptx_slide(prs, "内容依据 · 报名咨询")
    _pptx_body(end, prs, [b.citation_note, b.cta], size=16)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF 渲染（可打印宣传页）
# ---------------------------------------------------------------------------
def render_pdf(b: Brochure) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    ListFlowable, ListItem, HRFlowable)

    pdfmetrics.registerFont(UnicodeCIDFont(PDF_FONT))
    ink = colors.HexColor("#" + INK)
    muted = colors.HexColor("#" + MUTED)
    rule = colors.HexColor("#" + RULE)

    def st(name, **kw):
        base = dict(fontName=PDF_FONT, textColor=colors.black)
        base.update(kw)
        return ParagraphStyle(name, **base)

    title_s = st("t", fontSize=22, textColor=ink, leading=28, spaceAfter=4)
    sub_s = st("s", fontSize=11, textColor=muted, leading=16, spaceAfter=2)
    meta_s = st("m", fontSize=10, textColor=muted, leading=14, spaceAfter=10)
    h_s = st("h", fontSize=13, textColor=ink, leading=18,
             spaceBefore=10, spaceAfter=4)
    body_s = st("b", fontSize=10.5, leading=15, spaceAfter=4)
    li_s = st("li", fontSize=10.5, leading=15)
    note_s = st("n", fontSize=9, textColor=muted, leading=13)

    def section(title: str):
        flow.append(Paragraph(title, h_s))
        flow.append(HRFlowable(width="100%", thickness=0.4, color=rule, spaceAfter=6))

    def bullets(items):
        flow.append(ListFlowable(
            [ListItem(Paragraph(x, li_s), leftIndent=6) for x in items],
            bulletType="bullet", start="•", leftIndent=12))

    flow = []
    flow.append(Paragraph(b.title, title_s))
    flow.append(Paragraph(b.subtitle, sub_s))
    flow.append(Paragraph(b.meta_line, meta_s))
    flow.append(HRFlowable(width="100%", thickness=0.8, color=rule, spaceAfter=8))

    section("课程亮点")
    bullets(b.highlights)

    section("适合谁听")
    flow.append(Paragraph(f"本课程面向：<b>{b.audience}</b>。", body_s))
    if b.pain:
        flow.append(Paragraph("你将面对的问题：" + b.pain, body_s))
    if b.outcome:
        flow.append(Paragraph("学完你将能够：" + b.outcome, body_s))

    section("课程大纲")
    for m in b.outline:
        pts = ("　—　" + "；".join(m.points)) if m.points else ""
        flow.append(Paragraph(f"<b>模块{m.idx}　{m.title}</b>（{m.minutes} 分钟）{pts}", body_s))

    section("学习收获")
    bullets(b.takeaways)

    section("内容依据")
    flow.append(Paragraph(b.citation_note, note_s))

    section("报名与咨询")
    flow.append(Paragraph(b.cta, body_s))

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm, title=b.title)
    doc.build(flow)
    return out.getvalue()


# ---------------------------------------------------------------------------
# 打包
# ---------------------------------------------------------------------------
def build_materials_zip(b: Brochure, run_id: str) -> bytes:
    """把三件套打包成 zip 返回。文件名带 run_id 便于区分。"""
    safe = (run_id or "course").replace("/", "_")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{safe}-预热物料.docx", render_docx(b))
        z.writestr(f"{safe}-预热物料.pptx", render_pptx(b))
        z.writestr(f"{safe}-预热物料.pdf", render_pdf(b))
    return buf.getvalue()
